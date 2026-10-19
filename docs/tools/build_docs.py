"""Render docs from data.py:  python docs/tools/build_docs.py   (needs openpyxl)"""
import csv
import sys
from collections import OrderedDict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from data import BUGS, CASES, REQS  # noqa: E402

DOCS = Path(__file__).resolve().parents[1]
FONT = "Arial"

# ---------------------------------------------------------------- markdown / csv
def bug_report_md(b):
    detected = "\n".join(f"- {d}" for d in b["detected"])
    steps = "\n".join(f"{i}. {s}" for i, s in enumerate(b["steps"], 1))
    suite_hint = "UI bugs (BUG-011..BUG-015) are switched off with `VITE_SEEDED_BUGS=false` (rebuild the web container)." \
        if b["layer"] == "UI" else f"`FIXED_BUGS={b['id']}` (restart the api container) turns off only this bug."
    return f"""# {b['id']}: {b['title']}

| Field | Value |
|---|---|
| Issue type | Bug |
| Priority | {b['priority']} |
| Severity | {b['severity']} |
| Component | {b['component']} |
| Layer | {b['layer']} |
| Reference | {b['ref']} |
| Affects version | 1.0.0 (seeded build) |
| Environment | Docker Compose: Laravel 11 API + MySQL 8 + React 18; Chrome / Firefox / Edge / WebKit |
| Status | Open |

## Description
{b['title']}.

**Code location:** {b['where']}

## Steps to reproduce
{steps}

## Expected result
{b['expected']}

## Actual result
{b['actual']}

## Detected by
{detected}

## Suggested fix
{b['fix']}

## Retest
{suite_hint}
Re-run the tests listed under *Detected by*: they must pass, and the rest of the regression suite must stay green.

## Workflow history
| Date | From | To | By | Comment |
|---|---|---|---|---|
| | | Open | | Logged from failing automated test |
"""


def write_markdown_and_csv():
    out = DOCS / "bug-reports"
    out.mkdir(exist_ok=True)
    for b in BUGS:
        (out / f"{b['id']}.md").write_text(bug_report_md(b))

    rows = "\n".join(
        f"| [{b['id']}](bug-reports/{b['id']}.md) | {b['title']} | {b['layer']} | {b['severity']} | {b['priority']} | {b['ref']} |"
        for b in BUGS)
    (DOCS / "seeded-bugs.md").write_text(f"""# Seeded bug catalog

The app under test ships with {len(BUGS)} deliberate defects so the suites have something real to find.
Each bug has a JIRA-style report in [`bug-reports/`](bug-reports) and is detected by at least one automated test.

| ID | Summary | Layer | Severity | Priority | Reference |
|---|---|---|---|---|---|
{rows}

## Switches

| What | How |
|---|---|
| All API bugs (BUG-001..010) on / off | `SEEDED_BUGS=true` or `false` (env of the `api` container, or `.env` for docker compose) |
| Turn off individual API bugs (retest) | `FIXED_BUGS=BUG-003,BUG-004` |
| UI bugs (BUG-011..015) on / off | `VITE_SEEDED_BUGS=true` or `false` (set at build/start of the `web` container; compose passes `SEEDED_BUGS` to it) |

Typical demo: run the suites on the seeded build (red), log the failures in JIRA from the bug reports,
set `FIXED_BUGS` for one bug, re-run the same tests (green), close the ticket.
""")

    with open(DOCS / "jira-import.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["Summary", "Issue Type", "Priority", "Component/s", "Labels", "Labels", "Description"])
        for b in BUGS:
            desc = (f"*Steps to reproduce*\n" + "\n".join(f"# {s}" for s in b["steps"]) +
                    f"\n\n*Expected*\n{b['expected']}\n\n*Actual*\n{b['actual']}\n\n"
                    f"*Code location*\n{b['where']}\n\n*Detected by*\n" + "\n".join(f"* {d}" for d in b["detected"]) +
                    f"\n\n*Suggested fix*\n{b['fix']}")
            w.writerow([f"[{b['id']}] {b['title']}", "Bug", b["priority"], b["component"], "seeded-bug",
                        b["layer"].lower().replace(" / ", "-").replace(" ", "-"), desc])


# ---------------------------------------------------------------- xlsx
def write_xlsx():
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    HEAD_FILL = PatternFill("solid", start_color="1F3864")
    HEAD_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
    BODY = Font(name=FONT, size=10)
    INPUT = Font(name=FONT, size=10, color="0000FF")  # blue = cells you are meant to edit
    BOLD = Font(name=FONT, size=10, bold=True)
    thin = Side(style="thin", color="BFBFBF")
    BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
    WRAP = Alignment(wrap_text=True, vertical="top")

    def style_header(ws, row, ncols):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=row, column=c)
            cell.fill, cell.font, cell.border = HEAD_FILL, HEAD_FONT, BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="center")

    def widths(ws, ws_widths):
        for i, w in enumerate(ws_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    tc = wb.create_sheet("Test Cases")
    rq = wb.create_sheet("Requirements")
    bg = wb.create_sheet("Seeded Bugs")

    # ---- Test Cases
    heads = ["ID", "Module", "Req ID", "Title", "Type", "Layer", "Priority", "Steps", "Expected result",
             "Automated in", "Automation", "Seeded bug", "Status", "Notes"]
    tc.append(heads)
    style_header(tc, 1, len(heads))
    counters = OrderedDict()
    for (mod, req, title, typ, layer, prio, steps, expected, auto, bug) in CASES:
        counters[mod] = counters.get(mod, 0) + 1
        parts = [p.strip() for p in auto.split(";")]
        automation = "Manual" if parts == ["Manual"] else ("Partial" if "Manual" in parts else "Automated")
        tc.append([f"TC-{mod}-{counters[mod]:03d}", mod, req, title, typ, layer, prio, steps, expected, auto,
                   automation, bug or None, "Not Run", None])
    last = tc.max_row
    for row in tc.iter_rows(min_row=2, max_row=last):
        for cell in row:
            cell.font, cell.border, cell.alignment = BODY, BORDER, WRAP
        row[12].font = INPUT  # Status
        row[13].font = INPUT  # Notes
    widths(tc, [12, 8, 8, 44, 14, 8, 9, 48, 42, 24, 12, 11, 11, 28])
    tc.freeze_panes = "E2"
    tc.auto_filter.ref = f"A1:{get_column_letter(len(heads))}{last}"
    dv = DataValidation(type="list", formula1='"Not Run,Pass,Fail,Blocked"', allow_blank=False)
    tc.add_data_validation(dv)
    dv.add(f"M2:M{last + 200}")
    for text, color in (("Pass", "C6EFCE"), ("Fail", "FFC7CE"), ("Blocked", "FFEB9C")):
        tc.conditional_formatting.add(f"M2:M{last + 200}", CellIsRule(operator="equal", formula=[f'"{text}"'],
                                      fill=PatternFill("solid", start_color=color, end_color=color)))

    # ---- Requirements (formulas)
    rq.append(["Req ID", "Requirement", "Test cases", "Covered?", "Automated or partial"])
    style_header(rq, 1, 5)
    for i, (rid, name) in enumerate(REQS, start=2):
        rq.append([rid, name,
                   f"=COUNTIF('Test Cases'!$C$2:$C$1000,A{i})",
                   f'=IF(C{i}>0,"Yes","No")',
                   f"=COUNTIFS('Test Cases'!$C$2:$C$1000,A{i},'Test Cases'!$K$2:$K$1000,\"<>Manual\")"])
    for row in rq.iter_rows(min_row=2, max_row=rq.max_row):
        for cell in row:
            cell.font, cell.border, cell.alignment = BODY, BORDER, WRAP
    widths(rq, [9, 40, 12, 10, 20])
    rq.freeze_panes = "A2"

    # ---- Seeded bugs (formulas)
    bg.append(["Bug", "Summary", "Layer", "Severity", "Linked test cases", "Detected by"])
    style_header(bg, 1, 6)
    for i, b in enumerate(BUGS, start=2):
        bg.append([b["id"], b["title"], b["layer"], b["severity"],
                   f"=COUNTIF('Test Cases'!$L$2:$L$1000,A{i})", "; ".join(b["detected"])])
    for row in bg.iter_rows(min_row=2, max_row=bg.max_row):
        for cell in row:
            cell.font, cell.border, cell.alignment = BODY, BORDER, WRAP
    widths(bg, [10, 52, 10, 10, 16, 70])
    bg.freeze_panes = "A2"

    # ---- Summary (formulas)
    s = summary
    s["A1"], s["A1"].font = "Test case summary: QA Orders platform", Font(name=FONT, size=14, bold=True)
    s["A2"] = ("Legend: blue cells on 'Test Cases' (Status, Notes) are for you to edit; everything on this sheet is a formula "
               "and updates automatically. Example: set TC-AUTH-001 to Pass and watch 'Pass rate' change.")
    s["A2"].font, s["A2"].alignment = Font(name=FONT, size=9, italic=True), Alignment(wrap_text=True, vertical="top")
    s.merge_cells("A2:D2")
    s.row_dimensions[2].height = 40

    def block(top, title, items, col_letter, label_header="Item"):
        s.cell(row=top, column=1, value=title)
        s.cell(row=top, column=2, value="Count")
        s.cell(row=top, column=3, value="Share")
        style_header(s, top, 3)
        for k, item in enumerate(items, start=1):
            r = top + k
            s.cell(row=r, column=1, value=item).font = BODY
            s.cell(row=r, column=2, value=f"=COUNTIF('Test Cases'!${col_letter}$2:${col_letter}$1000,A{r})").font = BODY
            c = s.cell(row=r, column=3, value=f"=IFERROR(B{r}/$B$5,0)")
            c.font, c.number_format = BODY, "0.0%"
            for cc in range(1, 4):
                s.cell(row=r, column=cc).border = BORDER
        return top + len(items) + 2

    s["A4"], s["B4"] = "Headline metric", "Value"
    style_header(s, 4, 2)
    headline = [
        ("Total test cases", "=COUNTA('Test Cases'!$A$2:$A$1000)", "0"),
        ("Fully automated", "=COUNTIF('Test Cases'!$K$2:$K$1000,\"Automated\")", "0"),
        ("Partly automated", "=COUNTIF('Test Cases'!$K$2:$K$1000,\"Partial\")", "0"),
        ("Manual only", "=COUNTIF('Test Cases'!$K$2:$K$1000,\"Manual\")", "0"),
        ("Automated or partial (%)", "=IFERROR((B6+B7)/B5,0)", "0.0%"),
        ("Requirements covered (%)", "=IFERROR(COUNTIF(Requirements!$D$2:$D$21,\"Yes\")/COUNTA(Requirements!$A$2:$A$21),0)", "0.0%"),
        ("Seeded bugs with at least one linked case", "=COUNTIF('Seeded Bugs'!$E$2:$E$16,\">0\")", "0"),
        ("Executed (Pass + Fail + Blocked)", "=COUNTIF('Test Cases'!$M$2:$M$1000,\"Pass\")+COUNTIF('Test Cases'!$M$2:$M$1000,\"Fail\")+COUNTIF('Test Cases'!$M$2:$M$1000,\"Blocked\")", "0"),
        ("Pass rate of executed", "=IFERROR(COUNTIF('Test Cases'!$M$2:$M$1000,\"Pass\")/B12,0)", "0.0%"),
    ]
    for k, (label, formula, fmt) in enumerate(headline, start=5):
        s.cell(row=k, column=1, value=label).font = BODY
        c = s.cell(row=k, column=2, value=formula)
        c.font, c.number_format = BODY, fmt
        s.cell(row=k, column=1).border = s.cell(row=k, column=2).border = BORDER

    r = 5 + len(headline) + 1
    r = block(r, "By test type", ["API", "Functional", "Integration", "Regression", "Compatibility",
                                  "Security", "Accessibility", "Database"], "E")
    r = block(r, "By layer", ["API", "UI", "DB"], "F")
    r = block(r, "By priority", ["P1", "P2", "P3", "P4"], "G")
    r = block(r, "By module", ["AUTH", "ORD", "LIST", "FLOW", "STAT", "A11Y", "COMP", "DB", "CLOUD"], "B")
    r = block(r, "By status", ["Not Run", "Pass", "Fail", "Blocked"], "M")
    s.cell(row=r, column=1, value="Note: 'Share' columns are relative to the total number of test cases.").font = Font(name=FONT, size=9, italic=True)
    widths(s, [44, 14, 12, 12])

    wb.save(DOCS / "test-cases.xlsx")


if __name__ == "__main__":
    write_markdown_and_csv()
    write_xlsx()
    print("generated: seeded-bugs.md, bug-reports/*.md, jira-import.csv, test-cases.xlsx")
